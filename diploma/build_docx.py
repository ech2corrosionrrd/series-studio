#!/usr/bin/env python
"""Збирає всі файли diploma/NN-*.md в один Dyplom_Series_Lab.docx.

Оформлення (типові вимоги): Times New Roman 14 pt, інтервал 1,5, абзацний відступ 1,25 см,
поля: ліве 30, праве 15, верхнє й нижнє 20 мм, нумерація сторінок знизу по центру,
кожен заголовок рівня 1 починається з нової сторінки.
Схеми Mermaid підставляються з diploma/figures/ (PNG або JPG), у порядку появи в тексті.
Залежність: pip install python-docx pillow
"""
import glob
import io
import os
import re

from docx import Document
from PIL import Image
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Dyplom_Series_Lab.docx")
FIGURES = os.path.join(HERE, "figures")
FONT = "Times New Roman"

# Порядок збігається з появою ```mermaid у файлах 04 → 05 → 06 → 07 → 11.
MERMAID_FIGURES = [
    "rys-1-1",
    "rys-1-2",
    "rys-1-3",
    "rys-2-1",
    "rys-2-2",
    "rys-2-3",
    "rys-3-1",
    "rys-3-2",
    "rys-3-3",
    "rys-3-4",
    "rys-3-5",
    "rys-3-6",
    "rys-3-7",
    "rys-4-1",
    "rys-4-2",
    "rys-5-1",
    "rys-5-2",
    "dod-b1",
    "dod-b2",
    "dod-b3",
    "dod-b4",
]
MERMAID_CAPTIONS = [
    "Рисунок 1.1 — Контекстна діаграма Series Lab",
    "Рисунок 1.2 — Схема життєвого циклу створення та ончейн-випуску NFT-колекції",
    "Рисунок 1.3 — Дерево цілей і критеріїв ефективності Series Lab",
    "Рисунок 2.1 — Загальна архітектура системи",
    "Рисунок 2.2 — ER-діаграма основних таблиць бази даних",
    "Рисунок 2.3 — Блок-схема взаємодії користувача з Series Lab",
    "Рисунок 3.1 — Головна сторінка сайту Series Lab (опис етапів роботи)",
    "Рисунок 3.2 — Інтерфейс входу через Web3-гаманець",
    "Рисунок 3.3 — Етап 1. Огляд: параметри колекції",
    "Рисунок 3.4 — Етап 5. Генерація зображень",
    "Рисунок 3.5 — Етап 7. Експорт архіву випуску",
    "Рисунок 3.6 — Підготовка набору деплою",
    "Рисунок 3.7 — Діаграма послідовності оплати послуг токенами USDC",
    "Рисунок 4.1 — Схема конвеєра безперервної інтеграції (CI/CD)",
    "Рисунок 4.2 — Схема архітектури розгортання сервісу в Docker-контейнері",
    "Рисунок 5.1 — Структура вартості розробки Series Lab",
    "Рисунок 5.2 — Постійні витрати та точка беззбитковості",
    "Рисунок Б.1 — Діаграма варіантів використання",
    "Рисунок Б.2 — Діаграма активностей: генерація кадру",
    "Рисунок Б.3 — Стани замовлення на оплату",
    "Рисунок Б.4 — Шлях випуску",
]

doc = Document()
sec = doc.sections[0]
sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)


def set_font(style, size, bold=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


normal = doc.styles["Normal"]
set_font(normal, 14)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)

for lvl, size in ((1, 16), (2, 14), (3, 14), (4, 14)):
    st = doc.styles[f"Heading {lvl}"]
    set_font(st, size, bold=True)
    st.font.color.rgb = None
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_before = Pt(12 if lvl > 1 else 0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    run._r.append(el)

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")
mermaid_i = 0


def add_inline(par, text, size=None, force_bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Courier New"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            r.font.size = Pt(12 if size is None else min(size, 12))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        else:
            r = par.add_run(part)
        if force_bold:
            r.bold = True
        if size and not (part.startswith("`") and part.endswith("`")):
            r.font.size = Pt(size)


def shade(cell, color="D9D9D9"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]] if len(cells) > 2 else []
    ncols = len(header)
    t = doc.add_table(rows=1, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, h, size=12, force_bold=True)
        shade(c)
    for r in body:
        rc = t.add_row().cells
        for i in range(ncols):
            rc[i].text = ""
            p = rc[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, r[i] if i < len(r) else "", size=12)
    doc.add_paragraph()


def find_figure(stem):
    for ext in (".png", ".jpg", ".jpeg"):
        path = os.path.join(FIGURES, stem + ext)
        if os.path.isfile(path):
            return path
    return None


def figure_png_stream(path):
    """PNG у пам'яті — стабільніше відкривається в Word, ніж JPEG з mermaid.ink."""
    im = Image.open(path)
    im.load()
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
    else:
        im = im.convert("RGB")
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf, im.size


def picture_size_cm(pixel_w, pixel_h):
    max_w, max_h = Cm(16), Cm(22)
    if pixel_w <= 0 or pixel_h <= 0:
        return max_w
    w_cm = 16.0
    h_cm = w_cm * pixel_h / pixel_w
    if h_cm > 22.0:
        h_cm = 22.0
        w_cm = h_cm * pixel_w / pixel_h
    return Cm(w_cm)


def add_mermaid_figure():
    global mermaid_i
    if mermaid_i >= len(MERMAID_FIGURES):
        p = doc.add_paragraph()
        r = p.add_run("[Немає файлу рисунка для цієї схеми Mermaid]")
        r.italic = True
        return
    stem = MERMAID_FIGURES[mermaid_i]
    caption = MERMAID_CAPTIONS[mermaid_i]
    mermaid_i += 1
    path = find_figure(stem)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    if path:
        try:
            stream, (pw, ph) = figure_png_stream(path)
            run = p.add_run()
            run.add_picture(stream, width=picture_size_cm(pw, ph))
        except OSError as e:
            r = p.add_run(f"[Помилка рисунка {stem}: {e}]")
            r.italic = True
            r.font.size = Pt(11)
    else:
        r = p.add_run(f"[Немає файлу figures/{stem}.png|.jpg — згенеруйте з .mmd]")
        r.italic = True
        r.font.size = Pt(11)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    r = cp.add_run(caption)
    r.italic = True
    r.font.size = Pt(12)


def add_code(lines, lang=""):
    if lang == "mermaid":
        add_mermaid_figure()
        return
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(ln if ln else " ")
        r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        r.font.size = Pt(10)


def add_md(text, first_file):
    lines = text.split("\n")
    i = 0
    first_h1 = first_file
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            lang = s[3:].strip()
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip("\r"))
                i += 1
            add_code(buf, lang)
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(rows)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            if lvl == 1 and not first_h1:
                doc.add_page_break()
            first_h1 = False
            doc.add_heading(re.sub(r"[`*]", "", m.group(2)), level=lvl)
            i += 1
            continue
        if s in ("---", "***", ""):
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            add_inline(p, s.lstrip("> ").strip(), size=12)
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.25 + 0.6 * indent)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            marker = "– " if m.group(2) in "-*" else m.group(2) + " "
            p.add_run(marker)
            add_inline(p, m.group(3))
            i += 1
            continue
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#|```|\||>|\s*([-*]|\d+\.)\s)", lines[i]
        ):
            buf.append(lines[i].strip())
            i += 1
        # Підпис під рисунком уже додає збірка — пропускаємо дубль із Markdown.
        joined = " ".join(buf)
        if re.match(r"^\*\*Рисунок\s", joined) or re.match(r"^Рисунок\s+\d", joined):
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        add_inline(p, joined)


# Лише один розрив на розділ: перед кожним Heading 1 (крім першого).
# Окремий page break між .md-файлами давав порожні сторінки.
files = sorted(f for f in glob.glob(os.path.join(HERE, "[0-9][0-9]-*.md")))
for idx, f in enumerate(files):
    with open(f, encoding="utf-8") as fh:
        add_md(fh.read(), first_file=(idx == 0))

try:
    doc.save(OUT)
    saved = OUT
except PermissionError:
    alt = os.path.join(HERE, "Dyplom_Series_Lab_fresh.docx")
    doc.save(alt)
    saved = alt
    print("УВАГА: закрийте Word і перейменуйте або перезапустіть збірку.")
missing = [MERMAID_FIGURES[i] for i in range(mermaid_i, len(MERMAID_FIGURES))]
if mermaid_i != len(MERMAID_FIGURES):
    print("УВАГА: очікувалось", len(MERMAID_FIGURES), "блоків mermaid, оброблено", mermaid_i)
if missing:
    print("Немає файлів figures/ для:", ", ".join(missing))
print("OK:", saved, f"({len(files)} файлів, {mermaid_i} рисунків)")
